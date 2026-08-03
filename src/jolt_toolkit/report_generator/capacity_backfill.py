"""
capacity_backfill.py
====================
Backfill vehicles.json's ``effective_capacity_quarterly`` schema from
the existing xlsx report library, without re-running the reports.

For each EV (``fuel_type != "DIESEL"``):

1. Scan ``<report_db>/<REG>/jolt_report_<REG>_<start>_<end>.xlsx`` (matched
   strictly by file name, automatically excluding non-standard names such as
   ``*_finetuned*``).
2. Read the "Report" sheet, and for each report replicate the donor-based
   definition of :func:`_period_capacity_from_rows` (charge preferred: the mean
   ``Battery Capacity (kWh)`` of legs whose ``Energy Source`` is not
   ``soc_estimate`` and ``SOC Change (%) > 0``; otherwise discharge) to compute
   that period's ``(kwh, n_donors)``. This is the **same convention** as the
   generator / persistence path, ensuring backfill matches a re-run.
3. Write ``(kwh, n)`` into ``effective_capacity_quarterly[period_key]``, then
   recompute the donor weighted average from all reliable quarters
   (``n >= MIN_DONORS``) back into ``effective_capacity_kwh``, with sparse quarters'
   ``kwh`` backfilled to that average (see :func:`_recompute_weighted_capacity`).

Diesel vehicles are skipped entirely and get no quarterly fields. Pure
soc_estimate vehicles with no donor (e.g. the SOC-only Mercedes) produce an empty
quarterly and leave the existing scalar untouched.

Usage (jolt env, run from the repo root)::

    PYTHONUTF8=1 D:/Anaconda/envs/jolt/python.exe \\
        -m jolt_toolkit.report_generator.capacity_backfill \\
        --report-db excel_report_database/<version> [--dry-run]
"""

from __future__ import annotations

import argparse
import json
import logging
import re
from pathlib import Path

from filelock import FileLock
from openpyxl import load_workbook

from jolt_toolkit.configs import get_config_path
from jolt_toolkit.report_generator._generator import (
    MIN_DONORS,
    _period_capacity_from_rows,
    _recompute_weighted_capacity,
)

# Only match the standard report naming (ending exactly in _<8digit>_<8digit>.xlsx);
# finetuned / other suffixes naturally do not match and are skipped.
_REPORT_RE = re.compile(
    r"^jolt_report_(?P<reg>[A-Z0-9]+)_(?P<start>\d{8})_(?P<end>\d{8})\.xlsx$"
)


def _read_report_donor_capacity(xlsx_path: Path):
    """Read a single report's "Report" sheet, returning ``(kwh|None, n_donors, source)``.

    Uses ``data_only=True`` to read the cached values (Battery Capacity / SOC
    Change are plain numbers written by xlsxwriter, not ``=NA()`` formulas, so it
    is safe). Reuses :func:`_period_capacity_from_rows`: extracts each row into a
    ``(cap, soc, src)`` triple passed as idx 0/1/2, the same convention as the live path.
    """
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb["Report"] if "Report" in wb.sheetnames else wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = list(next(rows_iter))
        except StopIteration:
            return None, 0, "fallback"
        try:
            i_cap = header.index("Battery Capacity (kWh)")
            i_soc = header.index("SOC Change (%)")
            i_src = header.index("Energy Source")
        except ValueError:
            # Diesel / old column layouts do not have these columns
            return None, 0, "fallback"
        triples = []
        for r in rows_iter:
            if r is None:
                continue
            cap = r[i_cap] if i_cap < len(r) else None
            soc = r[i_soc] if i_soc < len(r) else None
            src = r[i_src] if i_src < len(r) else None
            triples.append((cap, soc, src))
    finally:
        wb.close()
    return _period_capacity_from_rows(triples, 0, 1, 2)


def backfill_vehicle(
    reg: str, report_dir: Path, entry: dict, min_donors: int = MIN_DONORS
) -> dict:
    """Backfill a single vehicle's quarterly + weighted average (mutates ``entry`` in place). Returns a summary dict."""
    quarterly: dict = {}
    per_period = (
        []
    )  # [(period_key, kwh|None, n, source)], includes fallback periods for display
    for xp in sorted(report_dir.glob(f"jolt_report_{reg}_*.xlsx")):
        m = _REPORT_RE.match(xp.name)
        if not m:
            continue  # skip non-standard names such as *_finetuned*
        period_key = f"{m.group('start')}_{m.group('end')}"
        kwh, n, src = _read_report_donor_capacity(xp)
        per_period.append((period_key, kwh, n, src))
        if src == "fallback" or kwh is None:
            continue  # periods with no donor do not enter quarterly
        quarterly[period_key] = {"kwh": round(float(kwh), 1), "n": int(n)}

    summary = {
        "reg": reg,
        "per_period": per_period,
        "old_kwh": entry.get("effective_capacity_kwh"),
        "quarterly": quarterly,
        "new_kwh": entry.get("effective_capacity_kwh"),
        "n_reliable": 0,
        "n_sparse": 0,
        "wrote": False,
    }
    if not quarterly:
        return summary

    wavg, n_rel, n_sparse = _recompute_weighted_capacity(quarterly, min_donors)
    entry["effective_capacity_quarterly"] = quarterly
    if wavg is not None:
        entry["effective_capacity_kwh"] = wavg
    summary.update(
        new_kwh=entry.get("effective_capacity_kwh"),
        n_reliable=n_rel,
        n_sparse=n_sparse,
        wrote=wavg is not None,
    )
    return summary


def _print_summary(s: dict, min_donors: int) -> None:
    reg = s["reg"]
    if not s["per_period"]:
        print(f"[no reports]  {reg}")
        return
    arrow = f"{s['old_kwh']} -> {s['new_kwh']}"
    print(
        f"\n== {reg} ==  effective_capacity_kwh: {arrow}  "
        f"(reliable={s['n_reliable']}, sparse={s['n_sparse']})"
    )
    for period_key, kwh, n, src in s["per_period"]:
        if src == "fallback" or kwh is None:
            print(f"    {period_key}  (no donor, src=fallback)  -- skipped")
            continue
        stored = s["quarterly"].get(period_key, {})
        tag = "SPARSE->avg" if n < min_donors else "reliable"
        print(
            f"    {period_key}  raw_kwh={round(float(kwh),1):>6}  n={n:>3}  "
            f"src={src:<9}  stored_kwh={stored.get('kwh'):>6}  [{tag}]"
        )


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Backfill effective_capacity_quarterly from existing xlsx reports"
    )
    ap.add_argument(
        "--report-db",
        required=True,
        help="report database version dir, e.g. excel_report_database/<version>",
    )
    ap.add_argument(
        "--min-donors",
        type=int,
        default=MIN_DONORS,
        help=f"reliability threshold on donor count (default {MIN_DONORS})",
    )
    ap.add_argument(
        "--dry-run",
        action="store_true",
        help="compute + print, but do NOT write vehicles.json",
    )
    args = ap.parse_args(argv)
    logging.basicConfig(level=logging.WARNING, format="%(message)s")

    db = Path(args.report_db)
    path = get_config_path("vehicles.json")
    # Guard the read-modify-write against a concurrent report-gen capacity
    # write-back (same lock file as _persist_effective_capacity).
    with FileLock(str(path) + ".lock"):
        with open(path, encoding="utf-8") as f:
            all_cfg = json.load(f)

        summaries = []
        for reg, cfg in all_cfg.items():
            if str(cfg.get("fuel_type", "")).upper() == "DIESEL":
                print(f"[skip diesel] {reg}")
                continue
            rdir = db / reg
            if not rdir.is_dir():
                print(f"[no reports]  {reg}")
                continue
            s = backfill_vehicle(reg, rdir, cfg, args.min_donors)
            summaries.append(s)
            _print_summary(s, args.min_donors)

        if args.dry_run:
            print("\n[dry-run] vehicles.json NOT written")
        else:
            with open(path, "w", encoding="utf-8") as f:
                json.dump(all_cfg, f, indent=2, ensure_ascii=False)
                f.write("\n")
            print(f"\nvehicles.json written: {path}")
    return summaries


if __name__ == "__main__":
    main()
