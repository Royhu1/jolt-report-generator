"""Report-sheet cell writing for a link column that holds NaN.

Only a non-empty string is a hyperlink. NaN is truthy, so a bare truthiness test
would hand it to xlsxwriter's ``write_url``; instead a NaN link cell must follow
the report's empty-cell convention and be written as the ``=NA()`` formula.
"""

from __future__ import annotations

from io import BytesIO

import numpy as np
import xlsxwriter
from openpyxl import load_workbook

from report_generator.columns import HEADERS
from report_generator.excel_writer import _write_report_sheet


def test_a_nan_link_is_written_as_the_na_formula():
    buffer = BytesIO()
    workbook = xlsxwriter.Workbook(buffer, {"in_memory": True})
    worksheet = workbook.add_worksheet("Report")
    row = [None] * (len(HEADERS) - 1)
    row[HEADERS.index("Leg Type") - 1] = "In Transit"
    row[HEADERS.index("Telematics Link") - 1] = np.nan

    _write_report_sheet(workbook, worksheet, [tuple(row)], HEADERS)
    workbook.close()

    reopened = load_workbook(BytesIO(buffer.getvalue()), data_only=False)
    link_cell = reopened["Report"].cell(
        row=2, column=HEADERS.index("Telematics Link") + 1
    )
    assert link_cell.value == "=NA()"
