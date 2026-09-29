"""The relabel patcher: an existing workbook's Leg Type column, in place.

Workbooks are written with the package's own writer from rows built inline, as
the generator builds them before its second pass (provisional labels: "In
Transit" trips and "<kind> Away" charges), Stop rows included. Every other cell
of every sheet, and the Report sheet's styles and hyperlinks, must come out as
they went in — except the Definitions sheet's Leg Type entry, which the patcher
brings up to date. The "old" glossaries are those of a report written before the
labels described the bases: an EV glossary without a Leg Type entry, a diesel one
whose first entry defines only "In Transit" and "Stop".
"""

from __future__ import annotations

import datetime
import json
import os

import openpyxl
import pandas as pd
import pytest

from report_generator import depots, leg_type_patcher
from report_generator.columns import DIESEL_HEADERS, HEADERS, _row_col_index
from report_generator.excel_writer import _write_excel_report
from report_generator.row_builder import _insert_stop_rows
from report_generator.xlsx_patch_common import save_workbook_atomically

NAN = float("nan")
#: The diesel glossary's Leg Type entry as reports were written before the labels
#: described the bases.
OLD_DIESEL_ENTRY = (
    'Leg Type: "In Transit" = trip (green); "Stop" = parked/idling gap between '
    "trips (white). Diesel vehicles have no charging events."
)
A = (52.0, -1.0)
B = (52.5, -1.5)
C = (52.2, -1.2)
T0 = pd.Timestamp("2026-03-02 00:00")


def _pt(p):
    return f"Point({p[0]:.6f} {p[1]:.6f})"


def _row(label, start, end, origin, dest, headers, *, distance=NAN, energy=NAN):
    r = [NAN] * (len(headers) - 1)

    def put(name, value):
        if name in headers:
            r[_row_col_index(name, headers)] = value

    put("Leg Type", label)
    put("Start Time (UTC)", start)
    put("End Time (UTC)", end)
    put("Origin (Lat, Lon)", _pt(origin))
    put("Destination (Lat, Lon)", _pt(dest))
    put("Duration (HH:MM:SS)", (end - start).total_seconds() / 86400.0)
    put("Distance (km)", distance)
    put("Operator", "OPX")
    put("Energy Change (kWh)", energy)
    put("Vehicle Mass (kg)", 31000.0)
    put("Energy Source", "lfc_fuel" if headers is DIESEL_HEADERS else "total_energy")
    put("Telematics Link", "https://data.example.org/explore/graphics/plots?x=1")
    put("SRF Logger Link", "https://data.example.org/explore/graphics/plots?l=1")
    return r


def _rows(headers, days=4, *, second_depot=False):
    """Days at A (A -> C -> A, a charge at A for EV); with ``second_depot``,
    a transfer to B and days based at B."""
    rows = []

    def at(d, hour):
        return T0 + pd.Timedelta(days=d, hours=hour)

    def day(d, base):
        rows.append(
            _row("In Transit", at(d, 7), at(d, 8), base, C, headers, distance=40.0)
        )
        rows.append(
            _row("In Transit", at(d, 12), at(d, 13), C, base, headers, distance=40.0)
        )
        if headers is HEADERS:
            rows.append(
                _row("DC Away", at(d, 14), at(d, 16), base, base, headers, energy=120.0)
            )

    for d in range(days):
        day(d, A)
    if second_depot:
        rows.append(
            _row("In Transit", at(days, 7), at(days, 9), A, B, headers, distance=65.0)
        )
        for d in range(days + 1, 2 * days + 1):
            day(d, B)
    return rows


def _write(tmp_path, rows, headers, name="jolt_report_TST01_20260302_20260310.xlsx"):
    path = tmp_path / name
    _write_excel_report(
        _insert_stop_rows(rows, headers=headers),
        "TST01",
        datetime.date(2026, 3, 2),
        datetime.date(2026, 3, 10),
        path,
        headers=headers,
    )
    return path


def _snapshot(path):
    """Every cell value of every sheet, and the Report sheet's look."""
    wb = openpyxl.load_workbook(path)
    values, looks = {}, {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                values[(ws.title, cell.coordinate)] = cell.value
                if ws.title == "Report":
                    looks[cell.coordinate] = (
                        cell.fill.fgColor.rgb,
                        cell.number_format,
                        cell.font.b,
                        cell.font.color.rgb if cell.font.color is not None else None,
                        cell.alignment.horizontal,
                        cell.border.left.style,
                        cell.hyperlink.target if cell.hyperlink is not None else None,
                    )
    meta = {
        "sheets": wb.sheetnames,
        "states": [ws.sheet_state for ws in wb.worksheets],
        "charts": [len(ws._charts) for ws in wb.worksheets],
        "freeze": wb["Report"].freeze_panes,
    }
    wb.close()
    return values, looks, meta


def _leg_type_column(path):
    ws = openpyxl.load_workbook(path)["Report"]
    return [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]


def _glossary(path):
    ws = openpyxl.load_workbook(path)["Definitions"]
    return [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]


def _age_glossary(path, headers):
    """Give a new workbook the glossary of a report written before the labels
    described the bases: the EV entry (the last) removed, the diesel entry (the
    first) set back to its old text."""
    wb = openpyxl.load_workbook(path)
    ws = wb["Definitions"]
    if headers is DIESEL_HEADERS:
        assert ws.cell(1, 1).value == depots.leg_type_definition(diesel=True)
        ws.cell(1, 1).value = OLD_DIESEL_ENTRY
    else:
        assert ws.cell(ws.max_row, 1).value == depots.leg_type_definition()
        ws.delete_rows(ws.max_row)
    wb.save(path)


@pytest.mark.parametrize("headers", [HEADERS, DIESEL_HEADERS], ids=["ev", "diesel"])
def test_only_the_leg_type_cells_change(tmp_path, headers):
    path = _write(tmp_path, _rows(headers, second_depot=True), headers)
    values0, looks0, meta0 = _snapshot(path)

    summary = leg_type_patcher.patch_workbook(path)
    values1, looks1, meta1 = _snapshot(path)

    assert summary["written"] is True and summary["changed"] > 0
    assert summary["layout"] == ("ev" if headers is HEADERS else "diesel")
    changed = {k for k in values0 if values0[k] != values1.get(k)}
    assert changed and all(
        sheet == "Report" and coord.startswith("B") for sheet, coord in changed
    )
    assert len(changed) == summary["changed"]
    assert set(values1) == set(values0)
    assert looks1 == looks0
    assert meta1 == meta0


def test_the_new_labels_are_the_depot_labels(tmp_path):
    path = _write(tmp_path, _rows(HEADERS, days=3), HEADERS)
    summary = leg_type_patcher.patch_workbook(path)
    stops = ["Stop"]
    day = ["Outbound", "Stop", "Return", "Stop", "DC Home"]
    assert _leg_type_column(path) == day + stops + day + stops + day
    (base,) = summary["bases"]["OPX"]
    assert (base["lat"], base["lon"], base["evidence"]) == (*A, "overnight")
    assert summary["changes"] == {
        "DC Away -> DC Home": 3,
        "In Transit -> Outbound": 3,
        "In Transit -> Return": 3,
    }


def test_a_second_run_changes_nothing_and_does_not_rewrite_the_file(tmp_path):
    path = _write(tmp_path, _rows(HEADERS, second_depot=True), HEADERS)
    leg_type_patcher.patch_workbook(path)
    before = path.read_bytes()
    mtime = path.stat().st_mtime_ns

    again = leg_type_patcher.patch_workbook(path)

    assert again["changed"] == 0 and again["written"] is False
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == mtime


def test_a_dry_run_reports_the_changes_and_writes_nothing(tmp_path):
    path = _write(tmp_path, _rows(HEADERS, second_depot=True), HEADERS)
    before = path.read_bytes()
    mtime = path.stat().st_mtime_ns

    summary = leg_type_patcher.patch_workbook(path, dry_run=True)

    assert summary["written"] is False
    assert summary["changed"] > 0 and summary["base_to_base"] == 1
    assert path.read_bytes() == before
    assert path.stat().st_mtime_ns == mtime
    assert list(tmp_path.iterdir()) == [path]  # no temporary file left behind


def test_a_workbook_labelled_by_the_generator_is_left_unchanged(tmp_path):
    rows = _rows(HEADERS, second_depot=True)
    depots.relabel_rows(rows)  # the generator's second pass
    path = _write(tmp_path, rows, HEADERS)
    before = path.read_bytes()

    summary = leg_type_patcher.patch_workbook(path)

    assert summary["changed"] == 0 and summary["written"] is False
    assert summary["definition"] == leg_type_patcher.DEFINITION_UNCHANGED
    assert path.read_bytes() == before


# ── The glossary follows the labels ──────────────────────────────────────────


def test_a_new_report_has_the_entry_last_for_ev_and_first_for_diesel(tmp_path):
    ev = _write(tmp_path, _rows(HEADERS, days=2), HEADERS, "jolt_report_EV_1_1.xlsx")
    dsl = _write(
        tmp_path,
        _rows(DIESEL_HEADERS, days=2),
        DIESEL_HEADERS,
        "jolt_report_D_1_1.xlsx",
    )
    assert _glossary(ev)[-1] == depots.leg_type_definition()
    assert _glossary(dsl)[0] == depots.leg_type_definition(diesel=True)
    assert sum(t.startswith("Leg Type:") for t in _glossary(ev)) == 1
    assert sum(t.startswith("Leg Type:") for t in _glossary(dsl)) == 1


@pytest.mark.parametrize("headers", [HEADERS, DIESEL_HEADERS], ids=["ev", "diesel"])
def test_an_old_glossary_is_brought_up_to_date_and_nothing_else_moves(
    tmp_path, headers
):
    path = _write(tmp_path, _rows(headers, second_depot=True), headers)
    new_glossary = _glossary(path)
    _age_glossary(path, headers)
    old_glossary = _glossary(path)
    values0, looks0, meta0 = _snapshot(path)

    summary = leg_type_patcher.patch_workbook(path)

    expected = (
        leg_type_patcher.DEFINITION_UPDATED
        if headers is DIESEL_HEADERS
        else leg_type_patcher.DEFINITION_ADDED
    )
    assert summary["definition"] == expected and summary["written"] is True
    # The glossary is now exactly a new report's: the one entry changed or
    # appended, every other entry where and as it was.
    assert _glossary(path) == new_glossary
    values1, looks1, meta1 = _snapshot(path)
    changed = {
        k for k in set(values0) | set(values1) if values0.get(k) != values1.get(k)
    }
    definitions = {k for k in changed if k[0] == "Definitions"}
    assert len(definitions) == 1
    ((_, coord),) = definitions
    if headers is DIESEL_HEADERS:
        assert coord == "A1" and values0[("Definitions", "A1")] == OLD_DIESEL_ENTRY
    else:
        assert coord == f"A{len(old_glossary) + 1}"
    assert all(
        sheet == "Report" and c.startswith("B") for sheet, c in changed - definitions
    )
    assert looks1 == looks0 and meta1 == meta0


def test_an_appended_entry_takes_the_style_of_the_entry_above(tmp_path):
    path = _write(tmp_path, _rows(HEADERS, days=2), HEADERS)
    _age_glossary(path, HEADERS)
    leg_type_patcher.patch_workbook(path)
    ws = openpyxl.load_workbook(path)["Definitions"]
    last, above = ws.cell(ws.max_row, 1), ws.cell(ws.max_row - 1, 1)
    assert last.value == depots.leg_type_definition()
    assert (last.alignment.wrap_text, last.alignment.vertical) == (True, "top")
    assert (last.alignment.wrap_text, last.alignment.vertical) == (
        above.alignment.wrap_text,
        above.alignment.vertical,
    )


@pytest.mark.parametrize("headers", [HEADERS, DIESEL_HEADERS], ids=["ev", "diesel"])
def test_a_second_run_leaves_a_brought_up_to_date_glossary_alone(tmp_path, headers):
    path = _write(tmp_path, _rows(headers, second_depot=True), headers)
    _age_glossary(path, headers)
    leg_type_patcher.patch_workbook(path)
    before = path.read_bytes()
    mtime = path.stat().st_mtime_ns

    again = leg_type_patcher.patch_workbook(path)

    assert again["definition"] == leg_type_patcher.DEFINITION_UNCHANGED
    assert again["changed"] == 0 and again["written"] is False
    assert path.read_bytes() == before and path.stat().st_mtime_ns == mtime


def test_the_glossary_is_brought_up_to_date_when_no_label_changes(tmp_path):
    rows = _rows(HEADERS, days=3)
    depots.relabel_rows(rows)  # labels already current
    path = _write(tmp_path, rows, HEADERS)
    _age_glossary(path, HEADERS)
    labels = _leg_type_column(path)

    summary = leg_type_patcher.patch_workbook(path)

    assert summary["changed"] == 0
    assert summary["definition"] == leg_type_patcher.DEFINITION_ADDED
    assert summary["written"] is True
    assert _leg_type_column(path) == labels
    assert _glossary(path)[-1] == depots.leg_type_definition()


def test_a_dry_run_reports_the_glossary_and_writes_nothing(tmp_path):
    path = _write(tmp_path, _rows(DIESEL_HEADERS, days=2), DIESEL_HEADERS)
    _age_glossary(path, DIESEL_HEADERS)
    before = path.read_bytes()

    summary = leg_type_patcher.patch_workbook(path, dry_run=True)

    assert summary["definition"] == leg_type_patcher.DEFINITION_UPDATED
    assert summary["written"] is False
    assert path.read_bytes() == before
    assert _glossary(path)[0] == OLD_DIESEL_ENTRY


def test_a_shorter_glossary_gets_the_entry_after_its_last_entry(tmp_path):
    # A report of an older generator, whose glossary had fewer entries, and a
    # sheet with a formatted but empty row below them.
    path = _write(tmp_path, _rows(HEADERS, days=2), HEADERS)
    wb = openpyxl.load_workbook(path)
    ws = wb["Definitions"]
    ws.delete_rows(4, ws.max_row - 3)
    ws.cell(6, 1).value = ""
    wb.save(path)

    summary = leg_type_patcher.patch_workbook(path)

    assert summary["definition"] == leg_type_patcher.DEFINITION_ADDED
    glossary = _glossary(path)
    assert glossary[3] == depots.leg_type_definition()
    assert all(not v for v in glossary[4:])


def test_an_ev_report_without_the_trailing_columns_is_relabelled(tmp_path):
    # A report written before the EP-confidence pair was appended.
    narrow = HEADERS[: HEADERS.index("EP Confidence")]
    rows = _rows(HEADERS, days=3)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(list(narrow))
    for k, r in enumerate(_insert_stop_rows(rows, headers=HEADERS), start=1):
        values = [k] + [None if (isinstance(v, float) and v != v) else v for v in r]
        ws.append(values[: len(narrow)])
    path = tmp_path / "jolt_report_TST02_20260302_20260305.xlsx"
    wb.save(path)

    summary = leg_type_patcher.patch_workbook(path)

    assert summary["layout"] == "ev"
    # No Definitions sheet to hold the entry: none is created.
    assert summary["definition"] == leg_type_patcher.DEFINITION_NO_SHEET
    assert openpyxl.load_workbook(path).sheetnames == ["Report"]
    assert _leg_type_column(path)[:5] == [
        "Outbound",
        "Stop",
        "Return",
        "Stop",
        "DC Home",
    ]


def test_a_workbook_the_labelling_cannot_read_is_refused_untouched(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(["Leg Number", "Leg Type", "Start Time (UTC)"])
    ws.append([1, "In Transit", datetime.datetime(2026, 3, 2, 7)])
    path = tmp_path / "jolt_report_TST03_20260302_20260302.xlsx"
    wb.save(path)
    before = path.read_bytes()
    with pytest.raises(ValueError, match="lack the column"):
        leg_type_patcher.patch_workbook(path)
    assert path.read_bytes() == before

    other = tmp_path / "jolt_report_TST04_20260302_20260302.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "Summary"
    wb.save(other)
    with pytest.raises(ValueError, match="no 'Report' sheet"):
        leg_type_patcher.patch_workbook(other)


# ── Finding the workbooks ────────────────────────────────────────────────────


def _touch_report(folder, name):
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / name
    path.write_bytes(b"")
    return path


def test_workbooks_are_collected_from_a_file_a_vehicle_or_a_tree(tmp_path):
    veh = tmp_path / "tree" / "TST01"
    a = _touch_report(veh, "jolt_report_TST01_20260101_20260331.xlsx")
    b = _touch_report(veh, "jolt_report_TST01_20260401_20260630.xlsx")
    _touch_report(veh, "jolt_report_TST01_20260101_20260331_finetuned.xlsx")
    _touch_report(veh, "~$jolt_report_TST01_20260401_20260630.xlsx")
    _touch_report(veh, "inspect_jolt_report_TST01_20260101_20260331.html")
    c = _touch_report(
        tmp_path / "tree" / "TST02", "jolt_report_TST02_20260101_20260331.xlsx"
    )
    (tmp_path / "tree" / "dashboard").mkdir()

    assert leg_type_patcher.collect_workbooks(a) == [a]
    assert leg_type_patcher.collect_workbooks(veh) == [a, b]
    assert leg_type_patcher.collect_workbooks(tmp_path / "tree") == [a, b, c]
    assert leg_type_patcher.collect_workbooks(tmp_path / "absent") == []


# ── The command line ─────────────────────────────────────────────────────────


def test_the_cli_dry_run_prints_the_summary_and_writes_only_the_json(tmp_path, capsys):
    folder = tmp_path / "TST01"
    folder.mkdir()
    path = _write(folder, _rows(HEADERS, second_depot=True), HEADERS)
    before = path.read_bytes()
    out_json = tmp_path / "summary.json"

    code = leg_type_patcher.main([str(folder), "--dry-run", "--json", str(out_json)])

    assert code == 0
    assert path.read_bytes() == before
    printed = capsys.readouterr().out
    assert "would change" in printed and "base-to-base" in printed
    assert "Leg Type definition unchanged" in printed
    assert "In Transit -> Outbound" in printed
    assert "bases, operator OPX: (52.0000, -1.0000) from overnight" in printed
    data = json.loads(out_json.read_text(encoding="utf-8"))
    assert data["dry_run"] is True
    (entry,) = data["workbooks"]
    assert entry["path"] == str(path) and entry["base_to_base"] == 1


def test_the_cli_relabels_in_place(tmp_path, capsys):
    path = _write(tmp_path, _rows(HEADERS, days=2), HEADERS)
    assert leg_type_patcher.main([str(path)]) == 0
    assert _leg_type_column(path)[0] == "Outbound"
    assert "changed 6" in capsys.readouterr().out  # 4 trips + 2 charges
    assert leg_type_patcher.main([str(path)]) == 0
    assert "changed 0" in capsys.readouterr().out


def test_the_cli_skips_a_workbook_open_in_excel(tmp_path):
    path = _write(tmp_path, _rows(HEADERS, days=2), HEADERS)
    (tmp_path / ("~$" + path.name)).write_bytes(b"lock")
    before = path.read_bytes()
    assert leg_type_patcher.main([str(path)]) == 1
    assert path.read_bytes() == before


def test_the_cli_reports_nothing_to_do(tmp_path):
    assert leg_type_patcher.main([str(tmp_path)]) == 1


# ── The atomic save ──────────────────────────────────────────────────────────


def test_a_failed_replace_leaves_the_workbook_and_no_temporary_file(
    tmp_path, monkeypatch
):
    path = _write(tmp_path, _rows(HEADERS, days=2), HEADERS)
    before = path.read_bytes()
    wb = openpyxl.load_workbook(path)
    wb["Report"].cell(2, 2).value = "Outbound"

    def refuse(_src, _dst):
        raise PermissionError("held by another process")

    monkeypatch.setattr(os, "replace", refuse)
    monkeypatch.setattr("report_generator.xlsx_patch_common._REPLACE_DELAY_S", 0.0)
    with pytest.raises(PermissionError):
        save_workbook_atomically(wb, path)
    assert path.read_bytes() == before
    assert list(tmp_path.iterdir()) == [path]


def test_an_atomic_save_replaces_the_workbook(tmp_path):
    path = _write(tmp_path, _rows(HEADERS, days=2), HEADERS)
    wb = openpyxl.load_workbook(path)
    wb["Report"].cell(2, 2).value = "Outbound"
    save_workbook_atomically(wb, path)
    assert _leg_type_column(path)[0] == "Outbound"
    assert list(tmp_path.iterdir()) == [path]
