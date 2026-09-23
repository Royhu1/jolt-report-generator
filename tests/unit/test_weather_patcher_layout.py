"""The coarse weather backfill over the two EV report widths and a diesel report.

``WeatherPatcher`` writes by fixed ``_COL_*`` index, so before it writes anything
it checks that the workbook really is the EV layout — a diesel report uses a
different column order and would silently get weather in the wrong cells.

A report tree is usually mixed: reports written before the two EP-confidence
columns were appended end at the column before ``EP Confidence``, newer ones
carry the pair. A check that demanded the full current ``HEADERS`` would refuse
every older report and the backfill would quietly patch nothing — returning 0 and
logging "diesel/unknown layout" for an ordinary EV report. These tests drive
``patch_file`` end to end over both widths, and over a diesel workbook, which must
still be refused without a byte changing.

No key and no network: the weather cache is stubbed so every location is already
cached, and the fetcher is removed so a fetch would fail loudly.
"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from report_generator.columns import DIESEL_HEADERS, HEADERS
from report_generator.weather_patcher import (
    _COL_DEST,
    _COL_END_TIME,
    _COL_LEG_TYPE,
    _COL_ORIGIN,
    _COL_START_TIME,
    _COL_TEMP,
    _COL_WEATHER_TYPE,
    _COL_WIND_DIR,
    WeatherPatcher,
)

# The header row of an EV report written before the EP-confidence pair existed.
PRE_EP_CONFIDENCE_HEADERS = HEADERS[: HEADERS.index("EP Confidence")]

# One canned observation: (temp, pressure, humidity, wind speed, wind dir, type).
_OBS = (7.5, 1011.0, 82.0, 4.0, 90.0, "Clouds")


class _StubKeys:
    """Stands in for the KeyManager, which ``patch_file`` only asks for a log line."""

    def summary(self):
        return {"active": 0, "total_keys": 0, "total_usage": 0}


class _StubCache:
    """Every location already cached, so no API key and no network are involved."""

    def get_batch(self, locs):
        return {loc: _OBS for loc in locs}, []

    def put_batch(self, fetched):  # pragma: no cover - never reached
        raise AssertionError("nothing should be missing from the stub cache")


def _patcher():
    """A ``WeatherPatcher`` with its cache stubbed and its fetcher removed.

    Built with ``__new__`` deliberately: ``__init__`` constructs a KeyManager,
    which wants an OpenWeather key this test has no business needing.
    """
    p = WeatherPatcher.__new__(WeatherPatcher)
    p._cache = _StubCache()
    p._keys = _StubKeys()
    p._fetcher = None  # a fetch would be a bug; let it fail loudly
    return p


def _workbook(tmp_path, headers, name):
    """One report workbook with a single trip row whose weather cells are empty."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(list(headers))
    ws.append([None] * len(headers))
    ws.cell(2, _COL_LEG_TYPE).value = "In Transit"
    # Naive, as a real report stores them (Excel carries no timezone, and the
    # patcher's _to_unix_utc reads a naive cell as UTC).
    ws.cell(2, _COL_START_TIME).value = datetime(2026, 1, 5, 8, 0)
    ws.cell(2, _COL_END_TIME).value = datetime(2026, 1, 5, 9, 0)
    ws.cell(2, _COL_ORIGIN).value = "Point(0.500000 0.500000)"
    ws.cell(2, _COL_DEST).value = "Point(0.600000 0.400000)"
    path = tmp_path / name
    wb.save(str(path))
    return path


def test_an_ev_report_without_the_ep_confidence_pair_is_patched(tmp_path):
    path = _workbook(tmp_path, PRE_EP_CONFIDENCE_HEADERS, "jolt_report_EVA_50col.xlsx")
    assert _patcher().patch_file(path) == 1

    ws = load_workbook(str(path))["Report"]
    assert ws.cell(2, _COL_TEMP).value == pytest.approx(7.5)
    assert ws.cell(2, _COL_WIND_DIR).value == "E"  # 90 deg
    assert ws.cell(2, _COL_WEATHER_TYPE).value == "Clouds"


def test_a_current_ev_report_is_patched(tmp_path):
    path = _workbook(tmp_path, HEADERS, "jolt_report_EVB_52col.xlsx")
    assert _patcher().patch_file(path) == 1
    ws = load_workbook(str(path))["Report"]
    assert ws.cell(2, _COL_TEMP).value == pytest.approx(7.5)


def test_a_diesel_report_is_skipped_without_writing_anything(tmp_path):
    path = _workbook(tmp_path, DIESEL_HEADERS, "jolt_report_DSLA_diesel.xlsx")
    before = path.read_bytes()
    assert _patcher().patch_file(path) == 0
    assert path.read_bytes() == before
