"""The fine-grained weather patcher samples every trip row along its track.

Which trips have a GPS track does not depend on their Leg Type: every trip row
is a window of the telematics feed. So a depot loop ("In House" / "Round Trip"),
an "Outbound" or a "Return" trip is sampled along its track exactly like an "In
Transit" one, and only a window without GPS samples falls back to the two
endpoints. Relabelling a report therefore never changes how its trips are
sampled.

No key and no network: the weather cache is stubbed so every location is
already cached, and a fetch would fail loudly.
"""

from __future__ import annotations

from datetime import datetime, timedelta

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from report_generator.columns import HEADERS
from report_generator.weather_fetcher.fine_grained_patcher import (
    FineGrainedWeatherPatcher,
    _resolve_col_indices,
)

COL = _resolve_col_indices(HEADERS)
DEPOT = (52.0, -1.0)
DAY = datetime(2026, 3, 2)


def _temperature(lat):
    """The stub's temperature: it varies with latitude, so the mean over a track
    differs from the mean of its two endpoints whenever the track leaves them."""
    return round((lat - 52.0) * 1000.0, 3)


class _StubCache:
    def get_batch(self, locs):
        return {
            loc: (_temperature(loc[0]), 1010.0, 80.0, 3.0, 90.0, "Clouds")
            for loc in locs
        }, []

    def put_batch(self, fetched):  # pragma: no cover - never reached
        raise AssertionError("nothing is missing from the stub cache")


class _StubFetcher:
    api_calls = 0
    failures = 0

    def reset_stats(self):
        pass

    def fetch_batch(self, *_a, **_k):  # pragma: no cover - never reached
        raise AssertionError("no location should need fetching")


class _StubKeys:
    def summary(self):
        return {"active": 0, "total_keys": 0, "total_usage": 0}


def _patcher():
    p = FineGrainedWeatherPatcher.__new__(FineGrainedWeatherPatcher)
    p._raw_dir = None
    p._raw_index = None
    p._min_interval = 60
    p._headers = HEADERS
    p._col_idx = COL
    p._cache = _StubCache()
    p._keys = _StubKeys()
    p._fetcher = _StubFetcher()
    p._stat_cache_hits = 0
    return p


def _loop_track(start):
    """A 20-minute loop from the depot 0.01 degree north and back, one sample a
    minute: (timestamp, lat, lon)."""
    lats = [52.0 + 0.001 * k for k in range(11)] + [
        52.0 + 0.001 * k for k in range(9, -1, -1)
    ]
    return [(start + timedelta(minutes=k), lat, DEPOT[1]) for k, lat in enumerate(lats)]


def _write(tmp_path, labels_and_starts, track_starts):
    """A report with one trip row per (label, start) — every one a loop starting
    and ending at the depot — and a raw_telematics day holding the loops that
    start at ``track_starts``."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append(list(HEADERS))
    for k, (label, start) in enumerate(labels_and_starts, start=2):
        ws.append([k - 1] + [None] * (len(HEADERS) - 1))
        ws.cell(k, COL["leg_type"]).value = label
        ws.cell(k, COL["start_time"]).value = start
        ws.cell(k, COL["end_time"]).value = start + timedelta(minutes=20)
        ws.cell(k, COL["origin"]).value = f"Point({DEPOT[0]:.6f} {DEPOT[1]:.6f})"
        ws.cell(k, COL["destination"]).value = f"Point({DEPOT[0]:.6f} {DEPOT[1]:.6f})"
    path = tmp_path / "jolt_report_TST01_20260302_20260302.xlsx"
    wb.save(path)

    raw = tmp_path / "raw_telematics"
    raw.mkdir()
    samples = [s for t in track_starts for s in _loop_track(t)]
    pd.DataFrame(
        {
            "eventDatetime": [f"{t.isoformat()}Z" for t, _, _ in samples],
            "latitude": [lat for _, lat, _ in samples],
            "longitude": [lon for _, _, lon in samples],
        }
    ).to_csv(raw / "raw_2026-03-02_0000.csv", index=False)
    return path


def _temperatures(path):
    ws = load_workbook(path)["Report"]
    return [ws.cell(r, COL["temp"]).value for r in range(2, ws.max_row + 1)]


TRACK_MEAN = round(
    float(np.mean([_temperature(lat) for _, lat, _ in _loop_track(DAY)])), 1
)
ENDPOINT_MEAN = _temperature(DEPOT[0])


def test_the_loop_is_a_real_test_of_the_sampling():
    assert TRACK_MEAN != pytest.approx(ENDPOINT_MEAN)


@pytest.mark.parametrize(
    "label", ["In Transit", "Round Trip", "In House", "Outbound", "Return"]
)
def test_every_trip_label_is_sampled_along_its_track(tmp_path, label):
    start = DAY.replace(hour=8)
    path = _write(tmp_path, [(label, start)], [start])
    stats = _patcher().patch_file(path, force_repatch=True)
    assert stats["patched_rows"] == 1
    assert stats["total_samples"] == 21  # the loop's samples, not two endpoints
    assert _temperatures(path) == [pytest.approx(TRACK_MEAN)]


def test_relabelling_a_trip_does_not_change_its_weather(tmp_path):
    starts = [DAY.replace(hour=8), DAY.replace(hour=10)]
    path = _write(
        tmp_path, [("In Transit", starts[0]), ("In House", starts[1])], starts
    )
    _patcher().patch_file(path, force_repatch=True)
    first, second = _temperatures(path)
    assert first == pytest.approx(second) == pytest.approx(TRACK_MEAN)


def test_a_window_without_samples_falls_back_to_the_endpoints(tmp_path):
    start = DAY.replace(hour=8)
    path = _write(tmp_path, [("Outbound", DAY.replace(hour=15))], [start])
    stats = _patcher().patch_file(path, force_repatch=True)
    assert stats["total_samples"] == 2
    assert _temperatures(path) == [pytest.approx(ENDPOINT_MEAN)]
