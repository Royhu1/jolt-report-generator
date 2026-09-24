"""Speed trips kept outside the capacity band (``keep_trips_outside_cap_band``).

``find_discharge_segments_by_speed`` drops a trip whose SOC-implied capacity
``|ΔE| / (|ΔSOC|/100)`` lies outside ``[cap_lo, cap_hi]``. On a feed whose energy
comes from a counter that leaves out what the battery spends while parked, while
the integer SOC includes it, short and medium trips imply too small a capacity
and real, speed-confirmed driving disappears from the report. With the
``speed_params`` key on, such a trip on counter energy is kept with
``effective_capacity_kwh`` ``None`` — so it is never a capacity donor — and a
private marker makes the mass split and merge and the anchor ordering give
nothing built from it a capacity either; ``run_segment_detection`` drops the
marker at the end. A trip on ``soc_estimate`` energy is dropped as before.

The frames are synthetic, one trip each, with hand-set counters: the energy and
ΔSOC below are exact, so every implied capacity is a hand computation.
"""

from __future__ import annotations

import copy
import datetime
import logging

import numpy as np
import openpyxl
import pandas as pd
import pytest

from report_generator import capacity_backfill
from report_generator._generator import JOLTReportGenerator
from report_generator.capacity import (
    _IDX_CAP,
    _IDX_ENERGY,
    _IDX_ESOURCE,
    _IDX_SOC_CHANGE,
    _IDX_START,
    _period_capacity_from_rows,
)
from report_generator.columns import HEADERS, _row_col_index
from report_generator.report_builder import _seg_to_row, _write_excel_report
from report_generator.segmentation import constants, detection
from report_generator.segmentation.constants import _CAPACITY_OUTSIDE_BAND_KEY
from report_generator.segmentation.detection import run_segment_detection
from report_generator.segmentation.mass_clustering import (
    _enforce_anchor_ordering,
    _merge_two_discharge_segs,
    _split_seg_at_times,
)
from report_generator.segmentation.speed_detection import (
    find_discharge_segments_by_speed,
)

TIME = constants.TIME_COL
SOC = constants.SOC_COL
ODO = constants.ODO_COL
SPEED = "wheel_based_speed"
MOV = "electric_energy_wheelbased_speed_over_zero"
TOT = "total_electric_energy_used_plugged_in_included"
MASS = "gross_combination_vehicle_weight"
DAY = "2026-08-12"
T0 = pd.Timestamp(f"{DAY} 08:00:00", tz="UTC")
# nominal 624 kWh -> the band is [312, 1248] kWh.
CAP_LO, CAP_HI = 312.0, 1248.0


def _trip_frame(
    energy_kwh: float,
    dsoc: float = 10.0,
    counters: tuple[str, ...] = (MOV,),
    stop_minutes: tuple[int, ...] = (),
    mass_kg: tuple[float, float] | None = None,
) -> pd.DataFrame:
    """One trip, 08:06 - 08:50 at 48 km/h, in a leg sampled every minute 08:00-09:00.

    The speed trip anchors to the zero-speed samples at 08:05 and 08:51 (the
    fleet's default ``zero_speed`` endpoints); between them the SOC falls by
    ``dsoc`` points exactly and each counter in ``counters`` rises by
    ``energy_kwh`` exactly. ``stop_minutes`` are standstill minutes inside the
    trip (bridged, being shorter than the stop gap). ``mass_kg`` gives the mass
    before and from 08:28 on, for a load change during a stop.
    """
    minutes = np.arange(61)
    moving = (minutes >= 6) & (minutes <= 50)
    speed = np.where(moving, 48.0, 0.0)
    for m in stop_minutes:
        speed[m] = 0.0
    progress = np.clip((minutes - 5) / 46.0, 0.0, 1.0)  # 0 at 08:05, 1 at 08:51
    frame = pd.DataFrame(
        {
            TIME: [(T0 + pd.Timedelta(minutes=int(m))).isoformat() for m in minutes],
            SOC: np.round(80.0 - dsoc * progress),
            SPEED: speed,
            ODO: 10_000.0 + np.cumsum(speed) / 60.0,
        }
    )
    for col in counters:
        frame[col] = 5_000_000.0 + energy_kwh * 1000.0 * progress
    if mass_kg is not None:
        frame[MASS] = np.where(minutes < 28, mass_kg[0], mass_kg[1])
    return frame.astype(str)


def _detect(frame, **kwargs):
    params = dict(
        speed_col=SPEED,
        speed_threshold_kmh=1.0,
        min_stop_duration_min=5.0,
        min_trip_duration_min=2.0,
        min_soc_drop=1.0,
        min_energy_kwh=1.0,
        cap_lo=CAP_LO,
        cap_hi=CAP_HI,
        total_energy_col=TOT,
        moving_energy_col=MOV,
        nominal_kwh=417.0,
    )
    params.update(kwargs)
    return find_discharge_segments_by_speed(frame, **params)


# ── The detector ─────────────────────────────────────────────────────────────


def test_by_default_a_trip_below_the_band_is_dropped():
    # 20 kWh over 10 SOC points implies 200 kWh, below the 312 kWh floor.
    assert _detect(_trip_frame(20.0)) == []


def test_with_the_key_a_counter_trip_below_the_band_is_kept_without_a_capacity():
    (trip,) = _detect(_trip_frame(20.0), keep_trips_outside_cap_band=True)
    assert trip["energy_source"] == "moving_energy"
    assert trip["delta_energy_kwh"] == pytest.approx(-20.0)
    assert trip["delta_soc_pct"] == pytest.approx(-10.0)
    assert trip["effective_capacity_kwh"] is None
    assert trip[_CAPACITY_OUTSIDE_BAND_KEY] is True
    # Everything else is measured as for any trip: 48 km/h for 45 minutes.
    assert trip["odo_end_km"] - trip["odo_start_km"] == pytest.approx(36.0)


def test_a_total_energy_trip_below_the_band_is_kept_as_well():
    frame = _trip_frame(20.0, counters=(TOT,))
    (trip,) = _detect(frame, keep_trips_outside_cap_band=True)
    assert trip["energy_source"] == "total_energy"
    assert trip["effective_capacity_kwh"] is None


def test_a_counter_trip_above_the_band_is_kept_without_a_capacity():
    # 150 kWh over 10 points implies 1500 kWh, above the 1248 kWh ceiling.
    assert _detect(_trip_frame(150.0)) == []
    (trip,) = _detect(_trip_frame(150.0), keep_trips_outside_cap_band=True)
    assert trip["effective_capacity_kwh"] is None


def test_a_soc_estimate_trip_outside_the_band_is_dropped_as_before():
    # No counter: the energy is 10 % x 200 kWh, so the implied capacity is the
    # 200 kWh used to estimate it, below the floor. Nothing measured the energy,
    # so there is nothing to keep.
    frame = _trip_frame(20.0, counters=())
    assert _detect(frame, nominal_kwh=200.0) == []
    assert _detect(frame, nominal_kwh=200.0, keep_trips_outside_cap_band=True) == []


def test_a_trip_inside_the_band_is_the_same_with_the_key(serialise):
    # 45 kWh over 10 points implies 450 kWh: in band, a donor like any other.
    frame = _trip_frame(45.0)
    off = _detect(frame)
    on = _detect(frame, keep_trips_outside_cap_band=True)
    assert serialise(on) == serialise(off)
    assert on[0]["effective_capacity_kwh"] == pytest.approx(450.0)
    assert _CAPACITY_OUTSIDE_BAND_KEY not in on[0]


@pytest.mark.parametrize("value", ["yes", 1, "true"])
def test_only_true_switches_the_key_on(value):
    assert _detect(_trip_frame(20.0), keep_trips_outside_cap_band=value) == []


def test_without_a_band_the_key_changes_nothing(serialise):
    frame = _trip_frame(20.0)
    off = _detect(frame, cap_lo=None, cap_hi=None)
    on = _detect(frame, cap_lo=None, cap_hi=None, keep_trips_outside_cap_band=True)
    assert serialise(on) == serialise(off)
    assert on[0]["effective_capacity_kwh"] == pytest.approx(200.0)


# ── Nothing built from a kept trip regains a capacity ────────────────────────


def _kept_trip(dsoc=20.0, energy_kwh=40.0):
    frame = _trip_frame(energy_kwh, dsoc=dsoc)
    (trip,) = _detect(frame, keep_trips_outside_cap_band=True)
    return trip, frame


def test_the_mass_split_gives_the_parts_of_a_kept_trip_no_capacity():
    trip, frame = _kept_trip()
    split_at = [pd.Timestamp(f"{DAY} 08:28:00", tz="UTC")]
    parts = _split_seg_at_times(trip, frame, split_at)
    assert len(parts) == 2
    # The energy is shared by SOC, so each part implies the trip's own 200 kWh.
    assert [p["delta_energy_kwh"] for p in parts] == pytest.approx([-20.0, -20.0])
    assert all(p["effective_capacity_kwh"] is None for p in parts)
    assert all(p[_CAPACITY_OUTSIDE_BAND_KEY] is True for p in parts)


def test_the_mass_split_of_an_ordinary_trip_is_unchanged():
    trip, frame = _kept_trip()
    ordinary = {k: v for k, v in trip.items() if k != _CAPACITY_OUTSIDE_BAND_KEY}
    ordinary["effective_capacity_kwh"] = 200.0
    split_at = [pd.Timestamp(f"{DAY} 08:28:00", tz="UTC")]
    parts = _split_seg_at_times(ordinary, frame, split_at)
    assert [p["effective_capacity_kwh"] for p in parts] == pytest.approx([200.0, 200.0])
    assert all(_CAPACITY_OUTSIDE_BAND_KEY not in p for p in parts)


def _segment(start, end, soc_s, soc_e, energy, cap, marked=False, anchors=None):
    seg = {
        "start_time": pd.Timestamp(f"{DAY} {start}", tz="UTC"),
        "end_time": pd.Timestamp(f"{DAY} {end}", tz="UTC"),
        "start_soc": soc_s,
        "end_soc": soc_e,
        "delta_soc_pct": soc_e - soc_s,
        "delta_energy_kwh": energy,
        "energy_source": "moving_energy",
        "delta_moving_kwh": None,
        "effective_capacity_kwh": cap,
        "odo_start_km": 0.0,
        "odo_end_km": 10.0,
        "lat_start": None,
        "lon_start": None,
        "lat_end": None,
        "lon_end": None,
        "_anchor_start_time": None,
        "_anchor_end_time": None,
        "_anchor_start_rel_kwh": float("nan"),
        "_anchor_end_rel_kwh": float("nan"),
    }
    if anchors is not None:
        (t_s, rel_s), (t_e, rel_e) = anchors
        seg["_anchor_start_time"] = pd.Timestamp(f"{DAY} {t_s}", tz="UTC")
        seg["_anchor_end_time"] = pd.Timestamp(f"{DAY} {t_e}", tz="UTC")
        seg["_anchor_start_rel_kwh"] = rel_s
        seg["_anchor_end_rel_kwh"] = rel_e
    if marked:
        seg[_CAPACITY_OUTSIDE_BAND_KEY] = True
    return seg


@pytest.mark.parametrize("kept_part", ["first", "second"])
def test_a_merge_that_takes_in_a_kept_trip_has_no_capacity(kept_part):
    # The kept part: 16 kWh over 8 points (200 kWh); the other: 45 kWh over 10
    # points (450 kWh). Merged, 61 kWh over 18 points would imply 338.9 kWh — in
    # the band — yet it rests on the kept part's SOC, so it carries no capacity.
    kept, other = ("first", "second") if kept_part == "first" else ("second", "first")
    parts = {
        kept: dict(energy=-16.0, soc=8.0, cap=None, marked=True),
        other: dict(energy=-45.0, soc=10.0, cap=450.0, marked=False),
    }
    first = _segment(
        "07:00",
        "07:40",
        90.0,
        90.0 - parts["first"]["soc"],
        parts["first"]["energy"],
        parts["first"]["cap"],
        marked=parts["first"]["marked"],
    )
    start = first["end_soc"]
    second = _segment(
        "08:00",
        "08:40",
        start,
        start - parts["second"]["soc"],
        parts["second"]["energy"],
        parts["second"]["cap"],
        marked=parts["second"]["marked"],
    )
    merged = _merge_two_discharge_segs(first, second)
    assert merged["delta_energy_kwh"] == pytest.approx(-61.0)
    assert merged["delta_soc_pct"] == pytest.approx(-18.0)
    assert merged["effective_capacity_kwh"] is None
    assert merged[_CAPACITY_OUTSIDE_BAND_KEY] is True


def test_a_merge_of_two_ordinary_trips_is_unchanged():
    merged = _merge_two_discharge_segs(
        _segment("07:00", "07:40", 90.0, 80.0, -45.0, 450.0),
        _segment("08:00", "08:40", 80.0, 70.0, -45.0, 450.0),
    )
    # 90 kWh over 20 points: 450 kWh.
    assert merged["effective_capacity_kwh"] == pytest.approx(450.0)
    assert _CAPACITY_OUTSIDE_BAND_KEY not in merged


def _overlapping_pair(marked: bool):
    # The first trip's end anchor (09:10, 30 kWh) overshoots the second trip's
    # start anchor (09:00, 20 kWh): the clamp brings its energy down to 20 kWh.
    cur = _segment(
        "08:00",
        "08:55",
        80.0,
        70.0,
        -30.0,
        None if marked else 300.0,
        marked=marked,
        anchors=(("07:58", 0.0), ("09:10", 30.0)),
    )
    nxt = _segment(
        "09:05",
        "10:00",
        70.0,
        60.0,
        -30.0,
        300.0,
        anchors=(("09:00", 20.0), ("10:05", 50.0)),
    )
    return cur, nxt


def test_the_anchor_clamp_leaves_a_kept_trip_without_a_capacity():
    cur, nxt = _overlapping_pair(marked=True)
    assert _enforce_anchor_ordering([cur, nxt], "UTCAP01") == 1
    assert cur["delta_energy_kwh"] == pytest.approx(-20.0)  # the energy is fixed
    assert cur["effective_capacity_kwh"] is None  # the capacity is not


def test_the_anchor_clamp_of_an_ordinary_trip_recomputes_its_capacity():
    cur, nxt = _overlapping_pair(marked=False)
    assert _enforce_anchor_ordering([cur, nxt], "UTCAP01") == 1
    # 20 kWh over 10 points: 200 kWh, as before this key existed.
    assert cur["effective_capacity_kwh"] == pytest.approx(200.0)


# ── Through run_segment_detection ────────────────────────────────────────────

REG = "UTCAP01"
PIPELINE = "ut_cap_speed"
_PIPELINE = {
    "branch": "speed",
    "charge_params": {"plateau_window_min": 60, "min_soc_rise": 5.0},
    "discharge_params": {
        "plateau_window_min": 15,
        "soc_rise_abort_pct": 3.0,
        "min_soc_drop": 5.0,
        "min_energy_kwh": 2.0,
    },
    "speed_params": {
        "speed_threshold_kmh": 1.0,
        "min_stop_duration_min": 5.0,
        "min_trip_duration_min": 2.0,
        "min_soc_drop": 1.0,
        "min_energy_kwh": 1.0,
    },
}
_VEHICLE = {
    "srf_reg": REG,
    "nominal_kwh": 624,
    "effective_capacity_kwh": 417.0,
    "pipeline": PIPELINE,
    "speed_col": SPEED,
    "moving_energy_col": MOV,
    "total_energy_col": TOT,
    "mass_col": MASS,
}


@pytest.fixture
def configure(monkeypatch):
    def _configure(keep=None, **vehicle):
        pipeline = copy.deepcopy(_PIPELINE)
        if keep is not None:
            pipeline["speed_params"]["keep_trips_outside_cap_band"] = keep
        monkeypatch.setitem(constants.PIPELINE_CONFIGS, PIPELINE, pipeline)
        monkeypatch.setitem(constants.VEHICLE_CONFIG, REG, {**_VEHICLE, **vehicle})

    return _configure


def _run(frame):
    return run_segment_detection(
        frame,
        reg=REG,
        suffix=f"{DAY}_0000",
        out_dir=None,
        generate_validation_fig=False,
        cap_lo=CAP_LO,
        cap_hi=CAP_HI,
    )


def test_the_pipeline_key_keeps_the_trip_and_the_marker_does_not_leave(
    configure, caplog
):
    configure(True)
    with caplog.at_level(logging.INFO, logger=detection.logger.name):
        _charges, (trip,) = _run(_trip_frame(20.0))
    assert trip["effective_capacity_kwh"] is None
    assert trip["delta_energy_kwh"] == pytest.approx(-20.0)
    assert _CAPACITY_OUTSIDE_BAND_KEY not in trip  # the public schema only
    assert "ep_audit" in trip
    assert any(
        "1 trips on counter energy kept" in r.getMessage() for r in caplog.records
    )


@pytest.mark.parametrize("keep", [None, False], ids=["absent", "false"])
def test_without_the_key_the_trip_is_dropped(configure, keep):
    configure(keep)
    _charges, discharges = _run(_trip_frame(20.0))
    assert discharges == []


def test_the_parts_of_a_kept_trip_split_at_a_load_change_carry_no_capacity(
    configure,
):
    # A three-minute stop at 08:26-08:28 (bridged: shorter than the 5-minute gap)
    # with the load changing from 20 t to 40 t: the mass split cuts the kept trip
    # at the first moving reading of the new load, 08:29.
    configure(True, split_by_mass=True, merge_by_mass=False)
    frame = _trip_frame(40.0, dsoc=20.0, stop_minutes=(26, 27, 28), mass_kg=(2e4, 4e4))
    _charges, parts = _run(frame)
    assert len(parts) == 2
    assert [p["delta_energy_kwh"] for p in parts] == pytest.approx([-20.0, -20.0])
    assert all(p["effective_capacity_kwh"] is None for p in parts)
    assert all(_CAPACITY_OUTSIDE_BAND_KEY not in p for p in parts)


# ── Downstream: the row, the capacity model, the workbook ────────────────────

I_EP = _row_col_index("Energy Performance (kWh/km)", HEADERS)
I_CONF = _row_col_index("EP Confidence", HEADERS)
I_DIST = _row_col_index("Distance (km)", HEADERS)


def _rows(configure):
    """A kept trip and an ordinary one, as report rows (the generator's order)."""
    configure(True)
    kept_frame = _trip_frame(20.0)
    _c, (kept,) = _run(kept_frame)
    donor_frame = _trip_frame(45.0)
    donor_frame[TIME] = [
        (pd.Timestamp(t) + pd.Timedelta(hours=3)).isoformat() for t in donor_frame[TIME]
    ]
    _c, (donor,) = _run(donor_frame)
    rows = []
    cumulative = 0.0
    for seg, frame in ((kept, kept_frame), (donor, donor_frame)):
        clean = {
            k: v for k, v in seg.items() if k not in constants._ANCHOR_PRIVATE_KEYS
        }
        row, cumulative = _seg_to_row(
            clean,
            "discharge",
            "https://data.example.org/api/legs/leg-1",
            [],
            [],
            frame,
            cumulative,
            None,
            srf_data=None,
            speed_col=SPEED,
            mass_col=MASS,
        )
        rows.append(list(row))
    return rows


def test_the_row_of_a_kept_trip_has_no_capacity_but_a_graded_ep(configure):
    kept, _donor = _rows(configure)
    assert kept[_IDX_CAP] is None
    assert kept[_IDX_ESOURCE] == "moving_energy"
    # EP from the counter energy over the odometer distance: 20 kWh / 36 km.
    assert kept[I_DIST] == pytest.approx(36.0)
    assert kept[I_EP] == pytest.approx(20.0 / 36.0, abs=1e-4)
    assert kept[I_CONF] in ("good", "caution", "poor")


def test_a_kept_trip_is_no_capacity_donor(configure):
    kept, donor = _rows(configure)
    assert donor[_IDX_CAP] == pytest.approx(450.0)
    # Only the ordinary trip counts: one donor of 450 kWh.
    assert _period_capacity_from_rows(
        [kept, donor], _IDX_CAP, _IDX_SOC_CHANGE, _IDX_ESOURCE
    ) == (pytest.approx(450.0), 1, "discharge")
    assert _period_capacity_from_rows(
        [kept], _IDX_CAP, _IDX_SOC_CHANGE, _IDX_ESOURCE
    ) == (None, 0, "fallback")


def test_the_generator_finalises_a_report_holding_a_kept_trip(configure):
    kept, donor = _rows(configure)
    gen = JOLTReportGenerator.__new__(JOLTReportGenerator)
    gen.debug_mode = False
    gen.fast_mode = True
    gen.srf_data = None
    out, period_kwh, period_n, period_src = gen._finalize_rows(
        [copy.copy(kept), copy.copy(donor)],
        HEADERS,
        is_diesel=False,
        cfg=dict(_VEHICLE),
        soc_est_cap=417.0,
        ep_audits={},
    )
    trips = [r for r in out if r[_row_col_index("Leg Type", HEADERS)] != "Stop"]
    kept_out = next(r for r in trips if r[_IDX_START] == kept[_IDX_START])
    # The capacity correction leaves the kept trip alone: no capacity, counter energy.
    assert kept_out[_IDX_CAP] is None
    assert kept_out[_IDX_ENERGY] == pytest.approx(-20.0)
    assert kept_out[I_CONF] in ("good", "caution", "poor")
    # The period capacity the ledger would record comes from the donor only.
    assert (period_kwh, period_n, period_src) == (
        pytest.approx(450.0),
        1,
        "discharge",
    )


def test_a_kept_trip_written_to_a_workbook_reads_back_as_no_donor(configure, tmp_path):
    kept, donor = _rows(configure)
    path = tmp_path / f"jolt_report_{REG}_20260812_20260813.xlsx"
    _write_excel_report(
        [kept, donor],
        REG,
        datetime.date(2026, 8, 12),
        datetime.date(2026, 8, 13),
        path,
        headers=HEADERS,
    )
    ws = openpyxl.load_workbook(path, data_only=True)["Report"]
    cap_col = HEADERS.index("Battery Capacity (kWh)") + 1
    assert ws.cell(2, cap_col).value in (None, "")  # the kept trip: blank
    assert ws.cell(3, cap_col).value == pytest.approx(450.0)
    # The ledger backfill reads the same donors back from the workbook.
    assert capacity_backfill._read_report_donor_capacity(path) == (
        pytest.approx(450.0),
        1,
        "discharge",
    )
