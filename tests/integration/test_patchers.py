"""Charger and Logger xlsx patchers, driven entirely offline.

Both patchers accept an injected ``srf_data`` AND pre-loaded windows / legs, so
the whole backfill can run without a client: ``ChargerPatcher.patch_file`` is
given ``charger_windows=`` and ``LoggerPatcher.patch_file`` is given
``logger_legs=`` / ``logger_windows=``. The logger legs are stand-ins backed by
the real DSL01 fixture, so the Channel-7 weather and CVW mass backfill run on
genuine 1 Hz readings.

Only ``patch_file`` is covered; the SRF-fetching ``_fetch_*`` paths are a
deliberate coverage gap (see tests/README.md).
"""

from __future__ import annotations

import datetime
from unittest.mock import Mock

import openpyxl
import pandas as pd
import pytest

from jolt_toolkit.report_generator.charger_patcher import (
    ChargerPatcher,
    _find_charger_matches,
    merge_save_charger_transactions,
)
from jolt_toolkit.report_generator.columns import HEADERS, _row_col_index
from jolt_toolkit.report_generator.logger_patcher import LoggerPatcher
from jolt_toolkit.report_generator.report_builder import _write_excel_report

# Windows inside the DSL01 fixture's 05:30:33 - 05:39:36 span.
CHARGE_START = pd.Timestamp("2025-10-07T05:31:00Z")
CHARGE_END = pd.Timestamp("2025-10-07T05:32:00Z")
TRIP_START = pd.Timestamp("2025-10-07T05:34:00Z")
TRIP_END = pd.Timestamp("2025-10-07T05:39:00Z")

CHARGER_URI = "https://data.example.org/api/transactions/tx-1"
LOGGER_URI = "https://data.example.org/api/legs/logger-1"


class FakeLoggerLeg:
    """A Logger leg stand-in serving slices of the real DSL01 fixture."""

    def __init__(self, frame, types=("7", "CVW")):
        self._frame = frame
        self.types = set(types)
        self.uri = LOGGER_URI
        self.start_time = frame.index[0]
        self.end_time = frame.index[-1]
        self.requested = []

    def get_data_frame(self, channel, resolution=None):
        self.requested.append(channel)
        if channel == "7":
            cols = [c for c in self._frame.columns if c.startswith("7 ")]
        elif channel == "CVW":
            cols = ["CVW gross combination vehicle weight"]
        else:  # pragma: no cover - the test only enables 7 / CVW
            return None
        return self._frame[cols].dropna(how="all")


def _row(leg_type, start, end, **overrides):
    row = [float("nan")] * (len(HEADERS) - 1)
    for link in ("Telematics Link", "Charger Link", "SRF Logger Link"):
        row[_row_col_index(link)] = None
    row[_row_col_index("Leg Type")] = leg_type
    row[_row_col_index("Start Time (UTC)")] = pd.Timestamp(start)
    row[_row_col_index("End Time (UTC)")] = pd.Timestamp(end)
    for header, value in overrides.items():
        row[_row_col_index(header)] = value
    return row


@pytest.fixture
def report_path(tmp_path):
    rows = [
        _row("DC Away", CHARGE_START, CHARGE_END),
        _row(
            "In Transit",
            TRIP_START,
            TRIP_END,
            **{
                "Distance (km)": 42.0,
                "Energy Change (kWh)": -50.0,
                "Elevation Difference (m)": 0.0,
            },
        ),
    ]
    out = tmp_path / "jolt_report_DSL01_20251007_20251008.xlsx"
    _write_excel_report(
        rows,
        "DSL01",
        datetime.date(2025, 10, 7),
        datetime.date(2025, 10, 8),
        out,
        headers=HEADERS,
    )
    return out


def _cells(path, header):
    ws = openpyxl.load_workbook(path)["Report"]
    col = HEADERS.index(header) + 1
    return [ws.cell(r, col) for r in range(2, ws.max_row + 1)]


# ── ChargerPatcher ───────────────────────────────────────────────────────────


def test_charger_patcher_accepts_an_injected_client():
    srf = Mock()
    patcher = ChargerPatcher(srf_data=srf)
    assert patcher._srf_data is srf  # no client was constructed


def test_charger_patcher_fills_the_link_and_the_energy(report_path):
    patched = ChargerPatcher(srf_data=Mock()).patch_file(
        report_path,
        charger_windows=[(CHARGE_START, CHARGE_END, CHARGER_URI, 55.5)],
    )
    assert patched == 1

    link_cells = _cells(report_path, "Charger Link")
    assert link_cells[0].value == "Link"
    assert link_cells[0].hyperlink.target.startswith(
        "https://data.example.org/explore/graphics/usage"
    )
    assert link_cells[1].value in (None, "")  # the trip row is untouched

    energy_cells = _cells(report_path, "Energy Output from Charger (kWh)")
    assert energy_cells[0].value == 55.5


def test_charger_patcher_sums_multiple_overlapping_transactions(report_path):
    windows = [
        (CHARGE_START, CHARGE_END, CHARGER_URI, 30.0),
        (CHARGE_START, CHARGE_END, CHARGER_URI + "-b", 25.0),
    ]
    ChargerPatcher(srf_data=Mock()).patch_file(report_path, charger_windows=windows)
    assert _cells(report_path, "Energy Output from Charger (kWh)")[0].value == 55.0


def test_charger_patcher_does_nothing_without_windows(report_path):
    assert (
        ChargerPatcher(srf_data=Mock()).patch_file(report_path, charger_windows=[]) == 0
    )
    assert _cells(report_path, "Charger Link")[0].value in (None, "")


def test_charger_patcher_ignores_a_non_overlapping_window(report_path):
    far = (
        pd.Timestamp("2025-10-08T05:31:00Z"),
        pd.Timestamp("2025-10-08T05:32:00Z"),
        CHARGER_URI,
        10.0,
    )
    assert (
        ChargerPatcher(srf_data=Mock()).patch_file(report_path, charger_windows=[far])
        == 0
    )


def test_charger_patcher_is_idempotent(report_path):
    windows = [(CHARGE_START, CHARGE_END, CHARGER_URI, 55.5)]
    patcher = ChargerPatcher(srf_data=Mock())
    assert patcher.patch_file(report_path, charger_windows=windows) == 1
    # A second pass finds the link already filled and rewrites nothing.
    assert patcher.patch_file(report_path, charger_windows=windows) == 0


def test_charger_patcher_missing_file_returns_zero(tmp_path):
    assert (
        ChargerPatcher(srf_data=Mock()).patch_file(
            tmp_path / "nope.xlsx", charger_windows=[]
        )
        == 0
    )


def test_find_charger_matches_tolerance_and_ordering():
    windows = [
        ("2025-10-07T06:00:00Z", "2025-10-07T06:30:00Z", "late", 1.0),
        ("2025-10-07T05:00:00Z", "2025-10-07T05:30:00Z", "early", 2.0),
    ]
    matches = _find_charger_matches(
        windows,
        pd.Timestamp("2025-10-07T05:29:00Z"),
        pd.Timestamp("2025-10-07T06:01:00Z"),
    )
    assert [m[2] for m in matches] == ["early", "late"]  # sorted by window start

    none_match = _find_charger_matches(
        windows,
        pd.Timestamp("2025-10-07T05:40:00Z"),
        pd.Timestamp("2025-10-07T05:50:00Z"),
        tol_min=4,
    )
    assert none_match == []


def test_merge_save_charger_transactions_round_trip(tmp_path):
    def _tx(uri, start, end, start_meter, end_meter):
        ct = Mock()
        ct.uri = uri
        ct.start_time = start
        ct.end_time = end
        ct.start_meter = start_meter
        ct.end_meter = end_meter
        ct.charger = Mock(
            label="L1", make="Nidec", model="DC-360", max_power=360, dc=True
        )
        return ct

    first = [_tx("tx-1", "2025-10-07T05:00:00Z", "2025-10-07T06:00:00Z", 10.0, 60.0)]
    assert merge_save_charger_transactions(first, tmp_path) == 1

    csv = tmp_path / "raw_charger" / "charger_transactions.csv"
    assert csv.exists()
    frame = pd.read_csv(csv)
    assert frame.loc[0, "energy_delivered_kwh"] == 50.0
    assert frame.loc[0, "charger_make"] == "Nidec"

    # A re-run with an overlapping uri de-duplicates rather than appending.
    second = [
        _tx("tx-1", "2025-10-07T05:00:00Z", "2025-10-07T06:00:00Z", 10.0, 70.0),
        _tx("tx-2", "2025-10-08T05:00:00Z", "2025-10-08T06:00:00Z", 0.0, 20.0),
    ]
    assert merge_save_charger_transactions(second, tmp_path) == 2
    frame = pd.read_csv(csv).set_index("uri")
    assert frame.loc["tx-1", "energy_delivered_kwh"] == 60.0  # freshest wins
    assert frame.loc["tx-2", "energy_delivered_kwh"] == 20.0


def test_merge_save_charger_transactions_no_objects(tmp_path):
    assert merge_save_charger_transactions([], tmp_path) == 0
    assert not (tmp_path / "raw_charger").exists()


# ── LoggerPatcher ────────────────────────────────────────────────────────────


def test_logger_patcher_accepts_an_injected_client():
    srf = Mock()
    assert LoggerPatcher(srf_data=srf)._srf_data is srf


def test_logger_patcher_fills_the_link_on_trip_rows_only(
    report_path, diesel_fixture_frame
):
    frame, _cfg = diesel_fixture_frame
    patched = LoggerPatcher(srf_data=Mock()).patch_file(
        report_path,
        logger_legs=[FakeLoggerLeg(frame)],
        logger_windows=[(TRIP_START, TRIP_END, LOGGER_URI)],
    )
    assert patched == 1

    link_cells = _cells(report_path, "SRF Logger Link")
    assert link_cells[0].value in (None, "")  # the charge row is skipped
    assert link_cells[1].value == "Link"
    assert link_cells[1].hyperlink.target.startswith(
        "https://data.example.org/explore/graphics/plots"
    )


def test_logger_patcher_backfills_channel_seven_weather(
    report_path, diesel_fixture_frame
):
    frame, _cfg = diesel_fixture_frame
    LoggerPatcher(srf_data=Mock()).patch_file(
        report_path,
        logger_legs=[FakeLoggerLeg(frame)],
        logger_windows=[(TRIP_START, TRIP_END, LOGGER_URI)],
    )
    assert _cells(report_path, "Average Temperature (C)")[1].value == pytest.approx(
        11.8, abs=0.3
    )
    assert _cells(report_path, "Average Pressure (hPa)")[1].value == 1024.0
    assert _cells(report_path, "Average Wind Speed (m/s)")[1].value == pytest.approx(
        2.1, abs=0.2
    )
    assert _cells(report_path, "Average Wind Direction")[1].value == "SW"


def test_logger_patcher_backfills_mass_from_the_cvw_channel(
    report_path, diesel_fixture_frame
):
    frame, _cfg = diesel_fixture_frame
    LoggerPatcher(srf_data=Mock()).patch_file(
        report_path,
        logger_legs=[FakeLoggerLeg(frame)],
        logger_windows=[(TRIP_START, TRIP_END, LOGGER_URI)],
    )
    assert _cells(report_path, "Vehicle Mass (kg)")[1].value == 16000


def test_logger_patcher_never_overwrites_an_existing_mass(
    tmp_path, diesel_fixture_frame
):
    frame, _cfg = diesel_fixture_frame
    rows = [
        _row(
            "In Transit",
            TRIP_START,
            TRIP_END,
            **{"Vehicle Mass (kg)": 31234.0, "Distance (km)": 42.0},
        )
    ]
    out = tmp_path / "jolt_report_DSL01_20251007_20251008.xlsx"
    _write_excel_report(
        rows,
        "DSL01",
        datetime.date(2025, 10, 7),
        datetime.date(2025, 10, 8),
        out,
        headers=HEADERS,
    )

    LoggerPatcher(srf_data=Mock()).patch_file(
        out,
        logger_legs=[FakeLoggerLeg(frame)],
        logger_windows=[(TRIP_START, TRIP_END, LOGGER_URI)],
    )
    assert _cells(out, "Vehicle Mass (kg)")[0].value == 31234.0


def test_logger_patcher_without_any_logger_data(report_path):
    assert (
        LoggerPatcher(srf_data=Mock()).patch_file(
            report_path, logger_legs=[], logger_windows=[]
        )
        == 0
    )


def test_logger_patcher_only_requests_the_channels_the_leg_advertises(
    report_path, diesel_fixture_frame
):
    frame, _cfg = diesel_fixture_frame
    leg = FakeLoggerLeg(frame, types=("7",))  # no CVW
    LoggerPatcher(srf_data=Mock()).patch_file(
        report_path,
        logger_legs=[leg],
        logger_windows=[(TRIP_START, TRIP_END, LOGGER_URI)],
    )
    assert set(leg.requested) == {"7"}
    # With no CVW the mass column stays at the writer's =NA() convention.
    assert _cells(report_path, "Vehicle Mass (kg)")[1].value in (None, "", "=NA()")


def test_logger_patcher_missing_file_returns_zero(tmp_path):
    assert (
        LoggerPatcher(srf_data=Mock()).patch_file(
            tmp_path / "nope.xlsx", logger_legs=[], logger_windows=[]
        )
        == 0
    )
