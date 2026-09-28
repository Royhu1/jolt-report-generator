"""Leg Types through the generator itself: the second pass runs on every report.

``JOLTReportGenerator.generate_report`` is driven over the committed fixtures
with the SRF surface replaced — one FPS leg (EV) or one SRF Logger leg (diesel)
whose data is the fixture, no charger transactions, fast mode. The rows are
built with provisional labels; once the run is segmented the generator finds
the run's bases and labels every trip and charge against them. What it writes
must therefore be exactly what the relabel patcher computes from the workbook's
own rows: patching a freshly generated report changes nothing.
"""

from __future__ import annotations

import copy
import json
import logging
import shutil
import types

import openpyxl
import pandas as pd
import pytest

from report_generator import _generator as gen_mod
from report_generator import leg_type_patcher
from report_generator.columns import _leg_is_charge, _leg_is_stop
from report_generator.data_class import ServerData
from report_generator.segmentation import constants

EV_ALIASES = ("EVSPD01", "EVSOC01", "EVMAD01", "EVSPD02")


@pytest.fixture
def offline_report(monkeypatch, tmp_path, frozen_configs, raw_fixture_path):
    """``run(alias) -> xlsx path``: one report of the fixture's day, offline."""
    cfg_dir = tmp_path / "configs"
    cfg_dir.mkdir()
    (cfg_dir / "vehicles.json").write_text(
        json.dumps(frozen_configs["vehicles"]), encoding="utf-8"
    )
    (cfg_dir / "pipelines.json").write_text(
        json.dumps(frozen_configs["pipelines"]), encoding="utf-8"
    )
    monkeypatch.setenv("JOLT_CONFIG_DIR", str(cfg_dir))
    monkeypatch.setenv("JOLT_CAPACITY_LEDGER", str(tmp_path / "ledger.json"))
    for alias, cfg in frozen_configs["vehicles"].items():
        monkeypatch.setitem(constants.VEHICLE_CONFIG, alias, copy.deepcopy(cfg))

    cache = tmp_path / "cache"
    (cache / "srf_raw").mkdir(parents=True)
    monkeypatch.setattr(gen_mod, "get_cache_dir", lambda: cache)

    def _no_vehicle(**_kwargs):
        raise LookupError("offline")

    client = types.SimpleNamespace(vehicles=types.SimpleNamespace(get=_no_vehicle))
    monkeypatch.setattr(gen_mod, "make_srf_client", lambda *a, **k: client)
    monkeypatch.setattr(
        gen_mod,
        "paging",
        types.SimpleNamespace(
            paged_items=lambda obj: list(obj) if isinstance(obj, list) else []
        ),
    )
    current = {}
    monkeypatch.setattr(
        gen_mod,
        "fetch_events",
        lambda **_kw: ServerData(
            vehicle=None, legs=[current["leg"]], charging_events=[]
        ),
    )

    def _trip(source):
        return types.SimpleNamespace(
            source=source,
            uri="trip-offline",
            trial=types.SimpleNamespace(description=None),
        )

    def run(alias):
        src = raw_fixture_path(alias)
        if alias.startswith("DSL"):
            frame = pd.read_csv(src, index_col=0)
            frame.index = pd.to_datetime(frame.index, utc=True)
            current["leg"] = types.SimpleNamespace(
                uri=f"https://data.example.org/api/legs/logger-{alias.lower()}",
                start_time=frame.index.min().to_pydatetime(),
                end_time=frame.index.max().to_pydatetime(),
                types=["CCVS", "LFC", "LFE", "VDHR", "CVW", "AMB", "7", "2"],
                get_data_frame=lambda *_a, **_k: frame,
                trip=_trip("SRFLOGGER_V1"),
            )
            day = frame.index.min().strftime("%Y-%m-%d")
        else:
            leg_id = f"offline-{alias.lower()}"
            shutil.copy2(src, cache / "srf_raw" / f"{leg_id}.csv")
            times = pd.to_datetime(
                pd.read_csv(src, dtype=str)["eventDatetime"], utc=True
            )
            current["leg"] = types.SimpleNamespace(
                uri=f"https://data.example.org/api/legs/{leg_id}",
                start_time=times.min().to_pydatetime(),
                end_time=times.max().to_pydatetime(),
                trip=_trip("FPS"),
            )
            day = times.min().strftime("%Y-%m-%d")
        generator = gen_mod.JOLTReportGenerator(
            report_output_folder=str(tmp_path / "out"), fast_mode=True
        )
        path = generator.generate_report(alias, day, day)
        assert path is not None, f"{alias}: no report written"
        return path

    return run


def _labels(path):
    ws = openpyxl.load_workbook(path)["Report"]
    return [ws.cell(r, 2).value for r in range(2, ws.max_row + 1)]


@pytest.mark.parametrize("alias", EV_ALIASES)
def test_an_ev_report_is_written_with_its_depot_labels(offline_report, alias):
    path = offline_report(alias)

    summary = leg_type_patcher.patch_workbook(path, dry_run=True)

    assert summary["changed"] == 0  # the generator already wrote these labels
    assert summary["definition"] == leg_type_patcher.DEFINITION_UNCHANGED
    assert summary["written"] is False
    assert summary["bases"]  # every EV fixture day yields a base
    trips = [
        lt for lt in _labels(path) if not _leg_is_charge(lt) and not _leg_is_stop(lt)
    ]
    assert {"Outbound", "Return"} <= set(trips)


def test_a_diesel_report_is_written_with_its_depot_labels(offline_report):
    path = offline_report("DSL01")
    summary = leg_type_patcher.patch_workbook(path, dry_run=True)
    assert summary["layout"] == "diesel"
    assert summary["changed"] == 0
    assert summary["definition"] == leg_type_patcher.DEFINITION_UNCHANGED
    assert summary["trips"] >= 1


def test_the_bases_are_logged(offline_report, caplog):
    with caplog.at_level(logging.INFO, logger="report_generator._generator"):
        offline_report("EVSOC01")
    messages = [r.getMessage() for r in caplog.records]
    assert any(m.startswith("Bases — operator") for m in messages)
    assert any(m.startswith("Leg types: ") for m in messages)


def test_a_labelling_failure_costs_the_labels_not_the_report(
    offline_report, monkeypatch, caplog
):
    def broken(*_a, **_k):
        raise RuntimeError("labelling broke")

    monkeypatch.setattr(gen_mod, "relabel_rows", broken)
    with caplog.at_level(logging.WARNING, logger="report_generator._generator"):
        path = offline_report("EVSOC01")
    # Provisional labels: trips "In Transit", charges "<kind> Away".
    labels = [lt for lt in _labels(path) if not _leg_is_stop(lt)]
    assert labels and all(lt == "In Transit" or lt.endswith(" Away") for lt in labels)
    assert any("Leg-type labelling failed" in r.getMessage() for r in caplog.records)
